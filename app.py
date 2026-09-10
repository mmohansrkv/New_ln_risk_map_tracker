"""
LN Risk Map — Productivity Tracker (Python / Flask version)
=============================================================
Same look, tabs and workflow as the original single-file HTML tool,
but rebuilt so the data lives on a server (data.json) instead of the
browser's localStorage. That means every employee and the admin see
the SAME shared data, from any device/browser, instead of each
browser having its own private copy.

Run:
    pip install -r requirements.txt
    python app.py
Then open:  http://127.0.0.1:5000/
"""

import json
import os
import threading
import uuid
from datetime import datetime, timedelta
from io import BytesIO

from flask import (Flask, jsonify, request, session, send_file,
                    render_template)
from openpyxl import Workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "data.json")

app = Flask(__name__)
app.secret_key = os.environ.get("LNRM_SECRET_KEY", "dev-secret-change-me")
app.permanent_session_lifetime = timedelta(hours=8)

_lock = threading.Lock()

DEFAULT_DB = {
    "admin": {"email": "admin@lnriskmap.com", "password": "Admin@123"},
    "employees": [],
    "processes": [],
    "productivity": [],
    "notes": [],
    "leaves": [],
}


# ----------------------------------------------------------------------
# Storage helpers
# ----------------------------------------------------------------------
def load_db():
    if not os.path.exists(DATA_FILE):
        save_db(DEFAULT_DB)
        return json.loads(json.dumps(DEFAULT_DB))
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        db = json.load(f)
    for key, val in DEFAULT_DB.items():
        db.setdefault(key, val if not isinstance(val, (list, dict)) else type(val)())
    return db


def save_db(db):
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(db, f, indent=2)


def new_id(prefix):
    return f"{prefix}_{uuid.uuid4().hex[:12]}"


# ----------------------------------------------------------------------
# Auth helpers
# ----------------------------------------------------------------------
def current_role():
    return session.get("role")


def current_emp_id():
    return session.get("empId")


def require_admin():
    return current_role() == "admin"


def find_employee(db, emp_id):
    return next((e for e in db["employees"] if e["empId"] == emp_id), None)


# ----------------------------------------------------------------------
# Page
# ----------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ----------------------------------------------------------------------
# Auth API
# ----------------------------------------------------------------------
@app.route("/api/login/admin", methods=["POST"])
def login_admin():
    data = request.get_json(force=True)
    email = (data.get("email") or "").strip()
    password = data.get("password") or ""
    db = load_db()
    if email == db["admin"]["email"] and password == db["admin"]["password"]:
        session.clear()
        session.permanent = True
        session["role"] = "admin"
        return jsonify({"ok": True, "role": "admin", "email": email})
    return jsonify({"ok": False, "error": "Invalid admin email or password."}), 401


@app.route("/api/login/user", methods=["POST"])
def login_user():
    data = request.get_json(force=True)
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""
    db = load_db()
    emp = next((e for e in db["employees"]
                if e["email"].lower() == email and e["password"] == password), None)
    if emp:
        session.clear()
        session.permanent = True
        session["role"] = "user"
        session["empId"] = emp["empId"]
        return jsonify({"ok": True, "role": "user", "employee": emp})
    return jsonify({"ok": False, "error": "Invalid email or password. Contact your admin."}), 401


@app.route("/api/logout", methods=["POST"])
def logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/me")
def me():
    role = current_role()
    if role == "admin":
        db = load_db()
        return jsonify({"role": "admin", "email": db["admin"]["email"]})
    if role == "user":
        db = load_db()
        emp = find_employee(db, current_emp_id())
        if emp:
            return jsonify({"role": "user", "employee": emp})
    session.clear()
    return jsonify({"role": None})


# ----------------------------------------------------------------------
# Employees (admin only)
# ----------------------------------------------------------------------
@app.route("/api/employees", methods=["GET"])
def list_employees():
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    return jsonify(db["employees"])


@app.route("/api/employees", methods=["POST"])
def save_employee():
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    edit_id = data.get("id")
    band = (data.get("band") or "").strip()
    emp_id = (data.get("empId") or "").strip()
    name = (data.get("name") or "").strip()
    email = (data.get("email") or "").strip()
    password = data.get("password") or ""

    if not (band and emp_id and name and email and password):
        return jsonify({"error": "All fields are required."}), 400

    with _lock:
        db = load_db()
        dup = next((e for e in db["employees"]
                    if e["email"].lower() == email.lower() and e["id"] != edit_id), None)
        if dup:
            return jsonify({"error": "Email already in use."}), 400

        if edit_id:
            emp = next((e for e in db["employees"] if e["id"] == edit_id), None)
            if not emp:
                return jsonify({"error": "Employee not found."}), 404
            old_emp_id, old_name, old_band = emp["empId"], emp["name"], emp["band"]
            emp.update({"band": band, "empId": emp_id, "name": name,
                        "email": email, "password": password})
            if old_emp_id != emp_id or old_name != name or old_band != band:
                for x in db["productivity"]:
                    if x["empId"] == old_emp_id:
                        x["empId"], x["empName"], x["band"] = emp_id, name, band
                for n in db["notes"]:
                    if n["empId"] == old_emp_id:
                        n["empId"] = emp_id
                for l in db["leaves"]:
                    if l["empId"] == old_emp_id:
                        l["empId"], l["empName"], l["band"] = emp_id, name, band
        else:
            db["employees"].append({
                "id": new_id("e"), "band": band, "empId": emp_id,
                "name": name, "email": email, "password": password,
            })
        save_db(db)
    return jsonify({"ok": True})


@app.route("/api/employees/<eid>", methods=["DELETE"])
def delete_employee(eid):
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    with _lock:
        db = load_db()
        db["employees"] = [e for e in db["employees"] if e["id"] != eid]
        save_db(db)
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Processes
# ----------------------------------------------------------------------
@app.route("/api/processes", methods=["GET"])
def list_processes():
    if not current_role():
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    return jsonify(db["processes"])


@app.route("/api/processes", methods=["POST"])
def save_process():
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    edit_id = data.get("id")
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Process name is required."}), 400
    target_hours = float(data.get("targetHours") or 0)
    target_100 = float(data.get("target100") or 0)
    target_count = float(data.get("targetCount") or 0)

    with _lock:
        db = load_db()
        if edit_id:
            p = next((p for p in db["processes"] if p["id"] == edit_id), None)
            if not p:
                return jsonify({"error": "Process not found."}), 404
            p.update({"name": name, "targetHours": target_hours,
                       "target100": target_100, "targetCount": target_count})
        else:
            db["processes"].append({
                "id": new_id("p"), "name": name, "targetHours": target_hours,
                "target100": target_100, "targetCount": target_count,
            })
        save_db(db)
    return jsonify({"ok": True})


@app.route("/api/processes/<pid>", methods=["DELETE"])
def delete_process(pid):
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    with _lock:
        db = load_db()
        db["processes"] = [p for p in db["processes"] if p["id"] != pid]
        save_db(db)
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Productivity
# ----------------------------------------------------------------------
@app.route("/api/productivity", methods=["GET"])
def list_productivity():
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    if role == "admin":
        emp_id = request.args.get("empId")
        rows = db["productivity"]
        if emp_id:
            rows = [x for x in rows if x["empId"] == emp_id]
        return jsonify(rows)
    return jsonify([x for x in db["productivity"] if x["empId"] == current_emp_id()])


@app.route("/api/productivity", methods=["POST"])
def submit_productivity():
    """Bulk-save process rows (+ optional note) for one date.
    Employees submit for themselves; admin submits on behalf of empId."""
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    date = data.get("date")
    entries = [r for r in (data.get("entries") or [])
               if r.get("processName") and (float(r.get("hour") or 0) > 0 or float(r.get("count") or 0) > 0)]
    note = data.get("note") or {}
    note_hour = float(note.get("hour") or 0)
    note_desc = (note.get("description") or "").strip()

    with _lock:
        db = load_db()
        if role == "admin":
            emp_id = data.get("empId")
            emp = find_employee(db, emp_id)
            if not emp:
                return jsonify({"error": "Please select an employee."}), 400
        else:
            emp = find_employee(db, current_emp_id())
            if not emp:
                return jsonify({"error": "Employee not found."}), 404

        if not date:
            return jsonify({"error": "Please select a date."}), 400
        if not entries:
            return jsonify({"error": "Add at least one process entry."}), 400

        total_hours = sum(float(r["hour"]) for r in entries)
        if total_hours + note_hour > 8:
            return jsonify({"error": "Total working hours (process + notes) cannot exceed 8 hours."}), 400

        for r in entries:
            db["productivity"].append({
                "id": new_id("pr"), "date": date, "band": emp["band"],
                "empId": emp["empId"], "empName": emp["name"],
                "processName": r["processName"],
                "hours": float(r["hour"]), "count": float(r.get("count") or 0),
            })

        if note_desc or note_hour > 0:
            db["notes"].append({
                "id": new_id("n"), "date": date, "empId": emp["empId"],
                "description": note_desc, "hour": note_hour,
            })

        save_db(db)
    return jsonify({"ok": True, "employee": emp["name"]})


@app.route("/api/productivity/<pid>", methods=["PUT"])
def edit_productivity(pid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    with _lock:
        db = load_db()
        rec = next((x for x in db["productivity"] if x["id"] == pid), None)
        if not rec:
            return jsonify({"error": "Not found."}), 404
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        if "hours" in data:
            rec["hours"] = float(data.get("hours") or 0)
        if "count" in data:
            rec["count"] = float(data.get("count") or 0)
        save_db(db)
    return jsonify({"ok": True})


@app.route("/api/productivity/<pid>", methods=["DELETE"])
def delete_productivity(pid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    with _lock:
        db = load_db()
        rec = next((x for x in db["productivity"] if x["id"] == pid), None)
        if not rec:
            return jsonify({"ok": True})
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        db["productivity"] = [x for x in db["productivity"] if x["id"] != pid]
        save_db(db)
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Notes
# ----------------------------------------------------------------------
@app.route("/api/notes", methods=["GET"])
def list_notes():
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    if role == "admin":
        emp_id = request.args.get("empId")
        rows = db["notes"]
        if emp_id:
            rows = [n for n in rows if n["empId"] == emp_id]
        return jsonify(rows)
    return jsonify([n for n in db["notes"] if n["empId"] == current_emp_id()])


@app.route("/api/notes/<nid>", methods=["PUT"])
def edit_note(nid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    with _lock:
        db = load_db()
        rec = next((n for n in db["notes"] if n["id"] == nid), None)
        if not rec:
            return jsonify({"error": "Not found."}), 404
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        if "description" in data:
            rec["description"] = (data.get("description") or "").strip()
        if "hour" in data:
            rec["hour"] = float(data.get("hour") or 0)
        save_db(db)
    return jsonify({"ok": True})


@app.route("/api/notes/<nid>", methods=["DELETE"])
def delete_note(nid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    with _lock:
        db = load_db()
        rec = next((n for n in db["notes"] if n["id"] == nid), None)
        if not rec:
            return jsonify({"ok": True})
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        db["notes"] = [n for n in db["notes"] if n["id"] != nid]
        save_db(db)
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Leaves
# ----------------------------------------------------------------------
@app.route("/api/leaves", methods=["GET"])
def list_leaves():
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    if role == "admin":
        emp_id = request.args.get("empId")
        rows = db["leaves"]
        if emp_id:
            rows = [l for l in rows if l["empId"] == emp_id]
        return jsonify(rows)
    return jsonify([l for l in db["leaves"] if l["empId"] == current_emp_id()])


@app.route("/api/leaves", methods=["POST"])
def add_leave():
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    date = data.get("date")
    reason = (data.get("reason") or "").strip()

    with _lock:
        db = load_db()
        if role == "admin":
            emp = find_employee(db, data.get("empId"))
            if not emp:
                return jsonify({"error": "Please select an employee."}), 400
            marked_by = "admin"
        else:
            emp = find_employee(db, current_emp_id())
            marked_by = "employee"

        if not date:
            return jsonify({"error": "Please select a leave date."}), 400
        if any(l["empId"] == emp["empId"] and l["date"] == date for l in db["leaves"]):
            return jsonify({"error": "A leave is already recorded for this employee on this date."}), 400

        db["leaves"].append({
            "id": new_id("lv"), "date": date, "empId": emp["empId"],
            "empName": emp["name"], "band": emp["band"],
            "reason": reason, "markedBy": marked_by,
        })
        save_db(db)
    return jsonify({"ok": True, "employee": emp["name"]})


@app.route("/api/leaves/<lid>", methods=["PUT"])
def edit_leave(lid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    data = request.get_json(force=True)
    with _lock:
        db = load_db()
        rec = next((l for l in db["leaves"] if l["id"] == lid), None)
        if not rec:
            return jsonify({"error": "Not found."}), 404
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        if "reason" in data:
            rec["reason"] = (data.get("reason") or "").strip()
        save_db(db)
    return jsonify({"ok": True})


@app.route("/api/leaves/<lid>", methods=["DELETE"])
def delete_leave(lid):
    role = current_role()
    if not role:
        return jsonify({"error": "forbidden"}), 403
    with _lock:
        db = load_db()
        rec = next((l for l in db["leaves"] if l["id"] == lid), None)
        if not rec:
            return jsonify({"ok": True})
        if role == "user" and rec["empId"] != current_emp_id():
            return jsonify({"error": "forbidden"}), 403
        db["leaves"] = [l for l in db["leaves"] if l["id"] != lid]
        save_db(db)
    return jsonify({"ok": True})


# ----------------------------------------------------------------------
# Excel export (admin only) — openpyxl, server-side
# ----------------------------------------------------------------------
def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 40)


def _sheet(wb, title, headers, rows):
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    _autosize(ws)
    return ws


def _xlsx_response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/export/<kind>")
def export_excel(kind):
    if not require_admin():
        return jsonify({"error": "forbidden"}), 403
    db = load_db()
    today = datetime.now().strftime("%Y-%m-%d")
    wb = Workbook()

    if kind == "productivity":
        rows = [[x["date"], x["empId"], x["empName"], x["band"], x["processName"], x["hours"], x["count"]]
                for x in db["productivity"]]
        _sheet(wb, "Productivity", ["Date", "EmpID", "Name", "Band", "Process", "Hours", "Count"], rows)
        return _xlsx_response(wb, f"all_productivity_{today}.xlsx")

    if kind == "notes":
        rows = [[n["date"], n["empId"], n["description"], n["hour"]] for n in db["notes"]]
        _sheet(wb, "Notes", ["Date", "EmpID", "Description", "Hour"], rows)
        return _xlsx_response(wb, f"all_notes_{today}.xlsx")

    if kind == "leaves":
        rows = [[l["date"], l["empId"], l["empName"], l["band"], l.get("reason", ""), l.get("markedBy", "")]
                for l in db["leaves"]]
        _sheet(wb, "Leaves", ["Date", "EmpID", "Name", "Band", "Reason", "MarkedBy"], rows)
        return _xlsx_response(wb, f"all_leaves_{today}.xlsx")

    if kind == "full":
        ws1 = wb.active
        ws1.title = "Employees"
        ws1.append(["EmpID", "Name", "Band", "Email"])
        for e in db["employees"]:
            ws1.append([e["empId"], e["name"], e["band"], e["email"]])
        _autosize(ws1)

        ws2 = wb.create_sheet("Processes")
        ws2.append(["Process", "TargetHours", "Target100", "TargetCountPerHour"])
        for p in db["processes"]:
            ws2.append([p["name"], p["targetHours"], p["target100"], p["targetCount"]])
        _autosize(ws2)

        ws3 = wb.create_sheet("Productivity")
        ws3.append(["Date", "EmpID", "Name", "Band", "Process", "Hours", "Count"])
        for x in db["productivity"]:
            ws3.append([x["date"], x["empId"], x["empName"], x["band"], x["processName"], x["hours"], x["count"]])
        _autosize(ws3)

        ws4 = wb.create_sheet("Notes")
        ws4.append(["Date", "EmpID", "Description", "Hour"])
        for n in db["notes"]:
            ws4.append([n["date"], n["empId"], n["description"], n["hour"]])
        _autosize(ws4)

        ws5 = wb.create_sheet("Leaves")
        ws5.append(["Date", "EmpID", "Name", "Band", "Reason", "MarkedBy"])
        for l in db["leaves"]:
            ws5.append([l["date"], l["empId"], l["empName"], l["band"], l.get("reason", ""), l.get("markedBy", "")])
        _autosize(ws5)

        return _xlsx_response(wb, f"LN_Risk_Map_full_database_{today}.xlsx")

    return jsonify({"error": "unknown export"}), 400


if __name__ == "__main__":
    load_db()  # ensure data.json exists on first run
    app.run(debug=True, host="0.0.0.0", port=5000)
